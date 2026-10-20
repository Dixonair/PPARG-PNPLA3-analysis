import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SETS = [
 ("GSEA_hallmark_human.csv", "Supplementary Table [X]",
  "Hallmark gene set enrichment in human subcutaneous white adipose tissue by PNPLA3 genotype.",
  "Genes were ranked by the DESeq2 Wald statistic for PNPLA3 GG (I148M) versus PNPLA3 CC; positive NES indicates higher expression in GG."),
 ("GSEA_hallmark_GSE162276.csv", "Supplementary Table [Y]",
  "Hallmark gene set enrichment in liver of hepatocyte-specific Pparg knockout mice on a high-fat diet (GSE162276).",
  "Genes were ranked by the DESeq2 Wald statistic for knockout versus control; positive NES indicates higher expression in the knockout."),
 ("GSEA_hallmark_GSE162276_TZDvsControl.csv", "Supplementary Table [Z]",
  "Hallmark gene set enrichment in liver of high-fat-diet mice treated with a thiazolidinedione (GSE162276).",
  "Genes were ranked by the DESeq2 Wald statistic for TZD versus control; positive NES indicates higher expression with TZD."),
]
NOTE = ("Pre-ranked GSEA (MSigDB v7.1 Hallmark collection, 50 sets) with 1,000 permutations. NES, normalised enrichment score; "
        "Nominal P, permutation P value (shown as < 0.001 when no permutation exceeded the observed score); FDR q, false discovery rate; "
        "Leading edge, genes in the leading edge / genes in the set found in the ranked list. Rows with FDR q < 0.05 are in bold.")

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcol); tcPr.append(s)
def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
def fmt_p(x): return "< 0.001" if x < 0.001 else f"{x:.3f}"
def name(t): 
    t = t.replace("HALLMARK_", "").replace("_", " ").title()
    for a, b in [("Tnfa", "TNFα"), ("Nfkb", "NF-κB"), ("Il2", "IL2"), ("Il6", "IL6"), ("Stat5", "STAT5"), ("Stat3", "STAT3"), ("Jak", "JAK"),
                 ("Kras", "KRAS"), ("Mtorc1", "mTORC1"), ("Myc", "MYC"), ("E2F", "E2F"), ("G2M", "G2M"), ("Dna", "DNA"), ("Uv ", "UV "),
                 ("Tgf Beta", "TGF-β"), ("P53", "p53"), ("Pi3K Akt Mtor", "PI3K/AKT/mTOR"), ("Wnt Beta Catenin", "WNT/β-catenin"), ("Il2 Stat5", "IL2/STAT5"),
                 ("Interferon Alpha", "Interferon α"), ("Interferon Gamma", "Interferon γ"), ("Reactive Oxygen", "Reactive oxygen"), ("Oxidative Phosphorylation", "Oxidative phosphorylation")]:
        t = t.replace(a, b)
    return t

doc = Document()
sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
for k, (f, label, title, rank) in enumerate(SETS):
    if k: doc.add_page_break()
    d = pd.read_csv(f); d["absnes"] = d["NES"].abs(); d = d.sort_values(["FDR q-val", "absnes"], ascending=[True, False]).reset_index(drop=True)
    p = doc.add_paragraph(); r = p.add_run(f"{label}. "); r.bold = True; p.add_run(title)
    p = doc.add_paragraph(rank); p.runs[0].italic = True
    t = doc.add_table(rows=1, cols=6); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["#", "Hallmark gene set", "NES", "Nominal P", "FDR q", "Leading edge"]
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(8.5)
        shade(c, "D9E2F3"); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
    repeat_header(t.rows[0])
    for i, row in d.iterrows():
        sig = row["FDR q-val"] < 0.05
        vals = [str(i + 1), name(row["Term"]), f"{row['NES']:+.2f}", fmt_p(row["NOM p-val"]), fmt_p(row["FDR q-val"]), str(row["Tag %"]).split()[0]]
        cells = t.add_row().cells
        for j, v in enumerate(vals):
            cells[j].text = ""; run = cells[j].paragraphs[0].add_run(v); run.font.size = Pt(8.5); run.bold = bool(sig)
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER
            if sig: shade(cells[j], "F2F2F2")
    t.autofit = False
    widths = [Cm(1.0), Cm(6.8), Cm(1.8), Cm(2.4), Cm(2.0), Cm(2.4)]
    for row in t.rows:
        for j, w in enumerate(widths): row.cells[j].width = w
    for j, w in enumerate(widths): t.columns[j].width = w
    n = doc.add_paragraph(NOTE); n.runs[0].font.size = Pt(8); n.paragraph_format.space_before = Pt(6)
doc.save("Supplementary_Hallmark_GSEA_tables.docx"); print("saved")
