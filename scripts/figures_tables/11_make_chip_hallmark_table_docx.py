import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcol); tcPr.append(s)
def fmt_p(x): return "< 0.001" if x < 0.001 else f"{x:.3f}"
def fmt_sci(x): return f"{x:.1e}".replace("e-0", "×10⁻").replace("e-", "×10⁻") if x < 0.001 else f"{x:.3f}"
def name(t):
    t = t.replace("HALLMARK_", "").replace("_", " ").title()
    for a, b in [("Tnfa", "TNFα"), ("Nfkb", "NF-κB"), ("Il2", "IL2"), ("Il6", "IL6"), ("Stat5", "STAT5"), ("Stat3", "STAT3"), ("Jak", "JAK"),
                 ("Kras", "KRAS"), ("Mtorc1", "mTORC1"), ("Myc", "MYC"), ("Dna", "DNA"), ("Uv ", "UV "), ("Tgf Beta", "TGF-β"), ("P53", "p53"),
                 ("Pi3K Akt Mtor", "PI3K/AKT/mTOR"), ("Wnt Beta Catenin", "WNT/β-catenin"), ("Interferon Alpha", "Interferon α"),
                 ("Interferon Gamma", "Interferon γ"), ("Oxidative Phosphorylation", "Oxidative phosphorylation"), ("E2F", "E2F"), ("G2M", "G2M")]:
        t = t.replace(a, b)
    return t
d = pd.read_csv("ChIP_bound_genes_Hallmark_ORA.csv").sort_values("padj").reset_index(drop=True)
N, nb = 16670, 3298
d["fold"] = d.k_bound_and_in_pathway / d.expected_if_random
doc = Document(); s = doc.sections[0]; s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
p = doc.add_paragraph(); r = p.add_run("Supplementary Table [X]. "); r.bold = True
p.add_run("Hallmark enrichment among genes bound by PPARG in human adipocytes (hASC-derived adipocytes, GSE115421).")
p = doc.add_paragraph(f"Genes with a consensus PPARG peak within ±5 kb of the TSS in at least 3 of 5 donors (n = {nb:,} genes expressed in the matched RNA-seq data) were tested against the MSigDB Hallmark sets with a one-sided hypergeometric test; the background was all {N:,} expressed genes.")
p.runs[0].italic = True
t = doc.add_table(rows=1, cols=8); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
hdr = ["#", "Hallmark gene set", "Genes in set (K)", "Bound genes in set (k)", "Expected by chance", "Fold enrichment", "P", "FDR q"]
widths = [Cm(0.9), Cm(4.9), Cm(1.6), Cm(1.9), Cm(1.8), Cm(2.2), Cm(2.0), Cm(2.0)]
for i, h in enumerate(hdr):
    c = t.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(8)
    shade(c, "D9E2F3"); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
for i, r_ in d.iterrows():
    sig = r_.padj < 0.05
    vals = [str(i + 1), name(r_.Term), str(int(r_.K_pathway_genes)), str(int(r_.k_bound_and_in_pathway)), f"{r_.expected_if_random:.1f}",
            f"{r_.fold:.2f}", fmt_sci(r_.pval), fmt_sci(r_.padj)]
    cells = t.add_row().cells
    for j, v in enumerate(vals):
        cells[j].text = ""; run = cells[j].paragraphs[0].add_run(v); run.font.size = Pt(8); run.bold = bool(sig)
        cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER
        if sig: shade(cells[j], "F2F2F2")
for row in t.rows:
    for j, w in enumerate(widths): row.cells[j].width = w
for j, w in enumerate(widths): t.columns[j].width = w
n = doc.add_paragraph(f"FDR q, Benjamini–Hochberg adjusted P across 50 sets; fold enrichment = bound genes in set / expected by chance. {int((d.padj<0.05).sum())} of 50 sets have FDR q < 0.05 (bold). "
                      "Because PPARG-bound genes are predominantly highly expressed in adipocytes and the background is restricted to expressed genes, "
                      "enrichment of many sets reflects general expression-related binding as well as adipocyte biology; the adipogenesis set shows the largest enrichment and the lowest P.")
n.runs[0].font.size = Pt(8); n.paragraph_format.space_before = Pt(6)
doc.save("Supplementary_ChIP_Hallmark_table.docx"); print("saved", int((d.padj<0.05).sum()))
