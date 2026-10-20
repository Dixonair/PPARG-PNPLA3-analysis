#!/usr/bin/env python3
"""
Source Data for the GSE115421 ChIP-seq panels:
    Main figure 6 panels e-l          (sheet names Fig6*)
    Supplementary Figure 7 panels a-f (SuppFig7*)
    Supplementary Figure 8 panels a,b (SuppFig8*)

EVERY value is read from the project files listed below or computed from them with the formulas stated in the
README sheet. Nothing is typed in or simulated. The 'Checks' sheet compares recomputed values with the numbers
printed on the figures (the figure values were read off the figure images and typed in below as FIG_LABELS).

Inputs (edit UP / OUT if your folders differ):
  UP  : PNPLA3_signal_profile.csv, FABP4_signal_profile.csv, all_peaks_annotated.csv, bound_genes_promoter.csv
  OUT : PNPLA3_/FABP4_quantitative_signal_metrics.csv, combined_8gene_heatmap_data.csv,
        ASCRosi_DE_with_ChIP_annotation.csv, ChIP_bound_genes_Hallmark_ORA.csv
Usage:  python3 make_source_data_GSE115421.py [UP] [OUT] [xlsx_path]
"""
import sys, os, hashlib
import numpy as np, pandas as pd
from scipy.stats import mannwhitneyu, fisher_exact

UP = sys.argv[1] if len(sys.argv) > 1 else "/mnt/user-data/uploads"
OUT = sys.argv[2] if len(sys.argv) > 2 else "/mnt/user-data/outputs"
XLSX = sys.argv[3] if len(sys.argv) > 3 else os.path.join(OUT, "GSE115421_PPARG", "Source_Data_Fig6e-l_SuppFig7_SuppFig8.xlsx")

SAMP = ["GSM3177950", "GSM3177951", "GSM3177952", "GSM3177953", "GSM3177954",
        "GSM3416379", "GSM3416380", "GSM3416381", "GSM3416382", "GSM3416383"]
GROUP = {s: ("hASC-derived adipocytes" if s.startswith("GSM31779") else "adipose tissue") for s in SAMP}
TSS = {"PNPLA3": 43923805, "FABP4": 81483350}               # as used by the figure scripts (08/09)
CHROM = {"PNPLA3": "chr22", "FABP4": "chr8"}
PEAK_PNPLA3 = (43925319, 43925915)                            # shaded intron-1 region in 09_rearranged_figure6.py
REPRO_GENES = ["ADIPOQ", "CD36", "FABP4", "PCK1", "PLIN1", "LPL", "SLC2A4", "PPARG", "PNPLA3"]
HEAT_ORDER = ["ADIPOQ", "CD36", "FABP4", "PLIN1", "PNPLA3", "MYOD1", "RHO", "INS"]

# numbers printed on the figures (typed from the images supplied by the author; used only for the Checks sheet)
FIG_LABELS = {
    "Fig6i  n not bound": 13372, "Fig6i  n PPARG-bound": 3298, "Fig6i  PNPLA3 t (1 dp)": 5.7,
    "Fig6i  PNPLA3 top % of bound genes": 8, "Fig6i  Mann-Whitney P (2 s.f.)": 8.2e-78,
    "Fig6j  OR (2 dp)": 5.80, "Fig6j  Fisher P (2 s.f.)": 1.3e-33, "Fig6j  % induced PPARG-bound (1 dp)": 3.6,
    "Fig6j  % induced not bound (1 dp)": 0.6,
    "SuppFig7a median width hASC": 241, "SuppFig7a median width tissue": 232,
    **{f"SuppFig7c peaks {s}": v for s, v in zip(SAMP, [21724, 39358, 24128, 41254, 32362, 5179, 36968, 4235, 28408, 13881])},
}

def rd(path): return pd.read_csv(path)
def md5(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for ch in iter(lambda: f.read(1 << 20), b""): h.update(ch)
    return h.hexdigest()

files = {
    "PNPLA3_signal_profile.csv": UP, "FABP4_signal_profile.csv": UP, "all_peaks_annotated.csv": UP,
    "bound_genes_promoter.csv": UP, "PNPLA3_quantitative_signal_metrics.csv": OUT,
    "FABP4_quantitative_signal_metrics.csv": OUT, "combined_8gene_heatmap_data.csv": OUT,
    "ASCRosi_DE_with_ChIP_annotation.csv": OUT, "ChIP_bound_genes_Hallmark_ORA.csv": OUT}
prof = {g: rd(f"{UP}/{g}_signal_profile.csv") for g in TSS}
met = {g: rd(f"{OUT}/{g}_quantitative_signal_metrics.csv") for g in TSS}
heat = rd(f"{OUT}/combined_8gene_heatmap_data.csv")
de = rd(f"{OUT}/ASCRosi_DE_with_ChIP_annotation.csv")
ora = rd(f"{OUT}/ChIP_bound_genes_Hallmark_ORA.csv")
prom = rd(f"{UP}/bound_genes_promoter.csv")
pk = rd(f"{UP}/all_peaks_annotated.csv")
pk = pk[pk.dataset.str.startswith("GSE115421")].copy()
pk["group"] = pk["sample"].map(GROUP)
sheets, checks = {}, []

def check(name, recomputed, fig_val, rel=None):
    ok = (abs(recomputed - fig_val) < 0.5 * 10 ** -rel) if rel is not None else bool(np.isclose(recomputed, fig_val, rtol=0.03))
    checks.append({"item": name, "recomputed_from_files": recomputed, "value_on_figure": fig_val, "match": "yes" if ok else "NO"})

# ------------------------------------------------------------------ tracks (Fig 6e,f; Supp Fig 8a,b)
def track_sheet(gene):
    d = prof[gene].copy()
    out = pd.DataFrame({"chromosome": CHROM[gene], "position_hg38_bin_start": d.position.astype(int),
                        "distance_from_TSS_bp": d.position.astype(int) - TSS[gene],
                        "distance_from_TSS_kb": (d.position - TSS[gene]) / 1000.0})
    for s in SAMP: out[s] = d[s].values
    return out
for gene, p6, p8 in [("PNPLA3", "Fig6e_PNPLA3_tracks", "SuppFig8a_PNPLA3_tracks"), ("FABP4", "Fig6f_FABP4_tracks", "SuppFig8b_FABP4_tracks")]:
    t = track_sheet(gene); sheets[p6] = t; sheets[p8] = t.copy()

# ------------------------------------------------------------------ Fig 6g occupancy
rows = []
for gene in ["FABP4", "PNPLA3"]:
    m = met[gene].copy()
    for _, r in m.iterrows():
        rows.append({"gene": gene, "sample": r["sample"], "group": GROUP[r["sample"]], "mean_signal_RPGC": r["mean_signal"]})
g_df = pd.DataFrame(rows)
summ = []
for (gene, grp), s in g_df.groupby(["gene", "group"], sort=False):
    v = s.mean_signal_RPGC.values
    summ.append({"gene": gene, "group": grp, "n": len(v), "mean": v.mean(), "SEM_ddof1": v.std(ddof=1) / np.sqrt(len(v))})
sheets["Fig6g_occupancy"] = g_df
sheets["Fig6g_summary"] = pd.DataFrame(summ)
# independent recomputation of the metric from the signal profile
win_rows = []
for gene in TSS:
    p = prof[gene]; w = p[(p.position >= TSS[gene] - 2000) & (p.position <= TSS[gene] + 2000)]
    diff = float((w[SAMP].mean() - met[gene].set_index("sample").loc[SAMP, "mean_signal"]).abs().max())
    win_rows.append({"gene": gene, "bins_in_window": len(w), "window_start_hg38": int(w.position.min()), "window_end_hg38": int(w.position.max()),
                     "window_start_rel_TSS_bp": int(w.position.min()) - TSS[gene], "window_end_rel_TSS_bp": int(w.position.max()) - TSS[gene],
                     "max_abs_diff_vs_metrics_csv": diff})
sheets["Fig6g_window_check"] = pd.DataFrame(win_rows)

# ------------------------------------------------------------------ Fig 6h reproducibility (+ Supp Fig 7 uses none of this)
cnt = prom[prom.SYMBOL.isin(REPRO_GENES) & prom.dataset.str.startswith("GSE115421")].pivot(index="SYMBOL", columns="dataset", values="n_samples_bound").reindex(REPRO_GENES)
cnt = cnt.rename(columns={"GSE115421_hASC": "hASC_samples_with_peak_of5", "GSE115421_tissue": "tissue_samples_with_peak_of5"}).reset_index().rename(columns={"SYMBOL": "gene"})
near = pk[pk.SYMBOL.isin(REPRO_GENES) & (pk.distanceToTSS.abs() <= 3000)]
ps_rows = []
for gene in REPRO_GENES:
    for s in SAMP:
        x = near[(near.SYMBOL == gene) & (near["sample"] == s)]
        if len(x):
            i = x.distanceToTSS.abs().idxmin()
            ps_rows.append({"gene": gene, "sample": s, "group": GROUP[s], "peak_within_3kb_of_TSS": "yes", "n_peaks_within_3kb": len(x),
                            "nearest_peak_distance_to_TSS_bp": int(x.loc[i, "distanceToTSS"]), "peak_chr": x.loc[i, "seqnames"],
                            "peak_start": int(x.loc[i, "start"]), "peak_end": int(x.loc[i, "end"])})
        else:
            ps_rows.append({"gene": gene, "sample": s, "group": GROUP[s], "peak_within_3kb_of_TSS": "no", "n_peaks_within_3kb": 0,
                            "nearest_peak_distance_to_TSS_bp": None, "peak_chr": None, "peak_start": None, "peak_end": None})
ps = pd.DataFrame(ps_rows)
chk = ps.groupby(["gene", "group"]).peak_within_3kb_of_TSS.apply(lambda s: (s == "yes").sum()).unstack()
cnt["recount_from_peaks_hASC"] = cnt.gene.map(chk["hASC-derived adipocytes"]); cnt["recount_from_peaks_tissue"] = cnt.gene.map(chk["adipose tissue"])
cnt["agrees_with_bound_genes_promoter_csv"] = np.where((cnt.hASC_samples_with_peak_of5 == cnt.recount_from_peaks_hASC) & (cnt.tissue_samples_with_peak_of5 == cnt.recount_from_peaks_tissue), "yes", "NO")
sheets["Fig6h_reproducibility"] = cnt
sheets["Fig6h_per_sample_peaks"] = ps

# ------------------------------------------------------------------ Fig 6i ECDF, Fig 6j enrichment
b = de[de.bound].t_stat.values; u = de[~de.bound].t_stat.values
mw = mannwhitneyu(b, u, alternative="greater")
pn = de[de.Gene == "PNPLA3"].iloc[0]; pct = (b < pn.t_stat).mean() * 100
sheets["Fig6i_ECDF_data"] = de[["Gene", "log2FC_Rosi_vs_DMSO", "t_stat", "pvalue", "padj", "bound", "induced"]].rename(columns={"Gene": "gene", "bound": "PPARG_bound_hASC", "induced": "induced_padj<0.05_and_log2FC>0"})
a = int(((de.bound) & (de.induced)).sum()); bb = int(((de.bound) & (~de.induced)).sum())
c = int(((~de.bound) & (de.induced)).sum()); d = int(((~de.bound) & (~de.induced)).sum())
orr, pf = fisher_exact([[a, bb], [c, d]], alternative="greater")
sheets["Fig6i_summary"] = pd.DataFrame([
    {"statistic": "n genes PPARG-bound (in RNA-seq table)", "value": len(b)}, {"statistic": "n genes not bound", "value": len(u)},
    {"statistic": "Mann-Whitney U, one-sided (bound > not bound), P", "value": mw.pvalue},
    {"statistic": "PNPLA3 paired t (Rosi vs DMSO)", "value": pn.t_stat}, {"statistic": "PNPLA3 log2FC", "value": pn.log2FC_Rosi_vs_DMSO},
    {"statistic": "PNPLA3 padj", "value": pn.padj},
    {"statistic": "PNPLA3 percentile rank among bound genes (% of bound genes with lower t)", "value": pct},
    {"statistic": "'top x%' printed on figure = 100 - percentile, rounded", "value": round(100 - pct)}])
sheets["Fig6j_enrichment"] = pd.DataFrame([
    {"group": "PPARG-bound", "induced": a, "not_induced": bb, "total": a + bb, "percent_induced": 100 * a / (a + bb)},
    {"group": "Not bound", "induced": c, "not_induced": d, "total": c + d, "percent_induced": 100 * c / (c + d)},
    {"group": "Fisher exact, one-sided (greater): odds ratio", "induced": orr, "not_induced": None, "total": None, "percent_induced": None},
    {"group": "Fisher exact, one-sided (greater): P", "induced": pf, "not_induced": None, "total": None, "percent_induced": None}])

# ------------------------------------------------------------------ Fig 6k,l heatmaps
M = heat.pivot(index="gene", columns="sample", values="mean_signal").loc[HEAT_ORDER, SAMP]
k = M.copy(); k.index.name = "gene"; sheets["Fig6k_raw_signal"] = k.reset_index()
kl = np.log2(M + 1); kl.index.name = "gene"
sheets["Fig6k_log2_plotted"] = kl.reset_index()
Z = (M.values - M.values.mean(1, keepdims=True)) / M.values.std(1, keepdims=True)   # population SD (ddof=0), as in 09_rearranged_figure6.py
zdf = pd.DataFrame(Z, index=HEAT_ORDER, columns=SAMP); zdf.index.name = "gene"; sheets["Fig6l_rowZ"] = zdf.reset_index()

# ------------------------------------------------------------------ Supp Fig 7
widths = pk[["dataset", "group", "sample", "width"]].rename(columns={"width": "peak_width_bp"})
sheets["SuppFig7a_peak_widths"] = widths
sheets["SuppFig7a_summary"] = (pk.groupby("group").width.agg(n_peaks="size", median_bp="median", min_bp="min", max_bp="max").reset_index())
pk["category"] = pk.annotation.str.replace(r"\s*\(.*", "", regex=True)
CATS = ["Promoter", "5' UTR", "Exon", "Intron", "3' UTR", "Downstream", "Distal Intergenic"]
cc = pk.groupby(["group", "category"]).size().unstack("group", fill_value=0).reindex(CATS, fill_value=0)
b7 = pd.DataFrame({"category": CATS})
for grp in ["hASC-derived adipocytes", "adipose tissue"]:
    b7[f"{grp}: n_peaks"] = cc[grp].values; b7[f"{grp}: percent_of_peaks"] = 100 * cc[grp].values / cc[grp].sum()
sheets["SuppFig7b_genomic_distribution"] = b7
sheets["SuppFig7c_peak_counts"] = pk.groupby("sample").size().reindex(SAMP).rename("n_peaks").reset_index().assign(group=lambda x: x["sample"].map(GROUP))[["sample", "group", "n_peaks"]]
top = ora.sort_values("padj").head(10).copy()
top["term_label"] = top.Term.str.replace("HALLMARK_", "").str.replace("_", " ").str.title()
top["minus_log10_padj"] = -np.log10(top.padj)
sheets["SuppFig7d_hallmark_ORA"] = top[["Term", "term_label", "K_pathway_genes", "k_bound_and_in_pathway", "expected_if_random", "pval", "padj", "minus_log10_padj", "genes"]]
for gene, nm in [("PNPLA3", "SuppFig7e_PNPLA3_group_mean"), ("FABP4", "SuppFig7f_FABP4_group_mean")]:
    p = prof[gene]
    sheets[nm] = pd.DataFrame({"chromosome": CHROM[gene], "position_hg38_bin_start": p.position.astype(int),
                               "hASC_mean_RPGC_n5": p[SAMP[:5]].mean(axis=1).values, "adipose_tissue_mean_RPGC_n5": p[SAMP[5:]].mean(axis=1).values})

# ------------------------------------------------------------------ checks against the numbers printed on the figures
check("Fig6i  n not bound", len(u), FIG_LABELS["Fig6i  n not bound"], 0); check("Fig6i  n PPARG-bound", len(b), FIG_LABELS["Fig6i  n PPARG-bound"], 0)
check("Fig6i  PNPLA3 t (1 dp)", pn.t_stat, FIG_LABELS["Fig6i  PNPLA3 t (1 dp)"], 1)
check("Fig6i  PNPLA3 top % of bound genes", round(100 - pct), FIG_LABELS["Fig6i  PNPLA3 top % of bound genes"], 0)
check("Fig6i  Mann-Whitney P (2 s.f.)", mw.pvalue, FIG_LABELS["Fig6i  Mann-Whitney P (2 s.f.)"])
check("Fig6j  OR (2 dp)", orr, FIG_LABELS["Fig6j  OR (2 dp)"], 2); check("Fig6j  Fisher P (2 s.f.)", pf, FIG_LABELS["Fig6j  Fisher P (2 s.f.)"])
check("Fig6j  % induced PPARG-bound (1 dp)", 100 * a / (a + bb), FIG_LABELS["Fig6j  % induced PPARG-bound (1 dp)"], 1)
check("Fig6j  % induced not bound (1 dp)", 100 * c / (c + d), FIG_LABELS["Fig6j  % induced not bound (1 dp)"], 1)
med = pk.groupby("group").width.median()
check("SuppFig7a median width hASC", med["hASC-derived adipocytes"], 241, 0); check("SuppFig7a median width tissue", med["adipose tissue"], 232, 0)
pc = pk.groupby("sample").size()
for s in SAMP: check(f"SuppFig7c peaks {s}", int(pc[s]), FIG_LABELS[f"SuppFig7c peaks {s}"], 0)
sheets["Checks"] = pd.DataFrame(checks)
n_bad = (sheets["Checks"].match != "yes").sum()

# ------------------------------------------------------------------ README
tpd = {g: f"{CHROM[g]}:{TSS[g]:,}" for g in TSS}
readme = [
 ("Purpose", "Source Data for Main Figure 6 panels e-l [confirm the figure number], Supplementary Figure 7 panels a-f and Supplementary Figure 8 panels a,b: PPARG ChIP-seq re-analysed from the raw reads of GSE115421 (hASC-derived adipocytes GSM3177950-54, n=5; adipose tissue GSM3416379-83, n=5)."),
 ("Provenance", "Each sheet holds the values read from, or computed with the formulas below from, the project files listed in the Files section. No value was typed in or simulated. The Checks sheet compares recomputed values with the numbers printed on the figures: %d of %d comparisons differ." % (n_bad, len(sheets["Checks"]))),
 ("Sheet naming", "Fig6e..Fig6l = main figure panels; SuppFig7a..f; SuppFig8a,b. Fig6e/f and SuppFig8a/b contain the same signal tracks (shown in different layouts); SuppFig7e/f show the group mean of those tracks."),
 ("Tracks (Fig6e,f; SuppFig7e,f; SuppFig8)", "Mean RPGC-normalised PPARG ChIP-seq signal in 20-bp bins from the bigWig files (04a_extract_locus_signal.py). position_hg38_bin_start = start of the 20-bp bin; distance_from_TSS = position - TSS used by the figure scripts (PNPLA3 %s; FABP4 %s). Group means (SuppFig7e,f) = arithmetic mean of the 5 samples per bin. Shaded intron-1 region in Supp Fig 8a (09_rearranged_figure6.py): chr22:%s-%s." % (tpd["PNPLA3"], tpd["FABP4"], f"{PEAK_PNPLA3[0]:,}", f"{PEAK_PNPLA3[1]:,}")),
 ("Fig6g", "mean_signal = mean of the 20-bp bin values in the window TSS-2 kb..TSS+2 kb (PNPLA3: 200 bins). See Fig6g_window_check: recomputing from the profile files reproduces the metrics CSVs (max abs difference in column max_abs_diff_vs_metrics_csv). Summary: mean and SEM (sample SD, n-1)."),
 ("!! FABP4 window", "The FABP4 signal profile starts at chr8:81,482,000, i.e. 1.35 kb upstream of the FABP4 'TSS' used (81,483,350), so the FABP4 window is TSS-1.35 kb..TSS+2 kb (168 of 200 bins), not TSS+/-2 kb. Panels g, k and l label it 'TSS+/-2 kb'. Either re-extract the FABP4 region from the bigWig files (set the region to start at 81,481,300 in 04a and re-run 04c) and regenerate, or change the label. The 'TSS' of FABP4 in the scripts is described as the promoter peak centre, not necessarily the annotated TSS - confirm."),
 ("Fig6h", "Number of samples (of 5 per group) with a peak within 3 kb of the gene TSS (bound_genes_promoter.csv, PROMOTER_WINDOW = 3,000 bp in 02_annotate_peaks.R). Fig6h_per_sample_peaks lists which samples and the nearest peak; the recount agrees with the CSV for every gene (column agrees_with_bound_genes_promoter_csv). Peaks were called without an input control; 'bound' needs >= 3 of 5 samples."),
 ("Fig6i,j", "ASCRosi_DE_with_ChIP_annotation.csv: rosiglitazone vs DMSO in hASC RNA-seq (script 05_integrate_with_rnaseq.py; patient-averaged paired t-test, NOT the limma model used for the primary hASC RNA-seq statistics). PPARG_bound = gene in the hASC promoter-bound list (>= 3 of 5 samples; 6,081 genes) AND present in the RNA-seq table (3,298 genes; the other genes of the list are not in the table). induced = padj < 0.05 and log2FC > 0. Mann-Whitney U is one-sided (bound > not bound); Fisher exact test one-sided. PNPLA3: bound, t = %.2f, log2FC = %.2f, padj = %.3f (not 'induced' by the padj < 0.05 criterion)." % (pn.t_stat, pn.log2FC_Rosi_vs_DMSO, pn.padj)),
 ("Fig6k,l", "mean_signal (TSS+/-2 kb mean RPGC; FABP4 window see above) per gene and sample from combined_8gene_heatmap_data.csv. Panel k shows log2(mean_signal + 1). Panel l = row z-score with the population SD (ddof = 0), as in 09_rearranged_figure6.py; the code that drew panel l in the final figure is not among the project files, so confirm that it used the same SD convention."),
 ("SuppFig7a", "Width of each peak (ChIPseeker table all_peaks_annotated.csv, GSE115421 only) per sample: 247,497 peaks = 158,826 hASC + 88,671 tissue. The histogram bin width used for the plotted figure is not recorded in the project files, so the unbinned widths are given; bin them as in your plotting code. Medians in SuppFig7a_summary."),
 ("SuppFig7b", "ChIPseeker annotation categories, detail text removed (e.g. 'Intron (ENST..., intron 2 of 9)' -> 'Intron'); peaks pooled across the 5 samples of each group; percent = n in category / n peaks of the group."),
 ("SuppFig7c", "Number of peaks per sample in all_peaks_annotated.csv (MACS2, q < 0.01, no control)."),
 ("SuppFig7d", "Over-representation analysis of PPARG-bound genes against the MSigDB Hallmark sets (ChIP_bound_genes_Hallmark_ORA.csv, 06_pathway_enrichment.py); the 10 terms with the lowest adjusted P are shown; minus_log10_padj = -log10(padj)."),
 ("Samples", "hASC-derived adipocytes: GSM3177950-54 (SRR7271363-67, 50 nt single-end); adipose tissue: GSM3416379-83 (SRR7963286-90, 100 nt single-end). Donor mapping of GSM to patient is not in the run table. Antibody: PPARg Santa Cruz sc-7196 (SRA run table). No input/IgG control in the series."),
 ("Reproducibility", "Regenerate this workbook with: python3 make_source_data_GSE115421.py [UP] [OUT] [xlsx]. The peak tables and tracks were produced by the earlier project scripts (01_reprocess_raw_reads.sh ... 08_generate_figures.py), which this workbook reads, not re-runs."),
]
files_tab = pd.DataFrame([{"file": f, "folder": d, "md5": md5(os.path.join(d, f))} for f, d in files.items()])
sheets_ordered = {"README": pd.DataFrame(readme, columns=["item", "description"]), "Files": files_tab}
sheets_ordered.update(sheets)

os.makedirs(os.path.dirname(XLSX), exist_ok=True)
with pd.ExcelWriter(XLSX, engine="openpyxl") as xw:
    for name, df in sheets_ordered.items():
        df.to_excel(xw, sheet_name=name[:31], index=False)
        ws = xw.sheets[name[:31]]
        from openpyxl.styles import Font
        for cell in ws[1]: cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, 1):
            w = min(max(len(str(col)), 12), 60) if name not in ("README",) else (28 if i == 1 else 140)
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        if name == "README":
            from openpyxl.styles import Alignment
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
print(f"wrote {XLSX}")
print(sheets["Checks"].to_string(index=False))
print("checks failing:", n_bad)
for n, df in sheets_ordered.items(): print(f"  {n[:31]:34s} {df.shape[0]:>7d} rows x {df.shape[1]} cols")
